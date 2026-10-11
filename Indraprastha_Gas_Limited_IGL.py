import os
import time
import re
import pandas as pd

from bs4 import BeautifulSoup
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC


# ============================================================
# CONFIGURATION
# ============================================================

BASE_URL = "https://www.iglonline.net"

OUTPUT_FILE = "IGL_CNG_Stations_All_Areas.xlsx"

WAIT_TIME = 15

PAGE_DELAY = 2


# ============================================================
# IGL AREAS
# ============================================================

AREAS = [
    ("Central Delhi", "central-delhi-cng-outlets"),
    ("East Delhi", "east-delhi-cng-outlets"),
    ("North Delhi", "north-delhi-cng-outlets"),
    ("North West Delhi", "north-west-delhi-cng-outlets"),
    ("South Delhi", "south-delhi-cng-outlets"),
    ("South West Delhi", "south-west-delhi-cng-outlets"),
    ("West Delhi", "west-delhi-cng-outlets"),
    ("Muzaffarnagar", "muzaffarnagar-cng-outlets"),
    ("Ghaziabad", "ghaziabad-cng-outlets"),
    ("Gautam Budh Nagar", "gautam-budh-nagar-cng-outlets"),
    ("Gurugram", "gurugram-cng-outlets"),
    ("Kanpur", "kanpur-cng-outlets"),
    ("Banda", "banda-cng-outlets"),
    ("Ajmer", "ajmer-cng-outlets"),
    ("Karnal", "karnal-cng-outlets"),
    ("Rewari", "rewari-cng-outlets"),
]


# ============================================================
# CLEAN TEXT
# ============================================================

def clean_text(value):

    if value is None:
        return ""

    value = str(value)

    value = value.replace("\xa0", " ")

    value = re.sub(
        r"\s+",
        " ",
        value
    )

    return value.strip()


# ============================================================
# CREATE CHROME DRIVER
# ============================================================

def create_driver():

    options = Options()

    # Normal Chrome window
    # Keep browser visible so you can see what is happening.
    # If everything works, you can uncomment headless mode.
    #
    # options.add_argument("--headless=new")

    options.add_argument(
        "--start-maximized"
    )

    options.add_argument(
        "--disable-blink-features=AutomationControlled"
    )

    options.add_argument(
        "--disable-notifications"
    )

    options.add_argument(
        "--disable-popup-blocking"
    )

    options.add_argument(
        "--no-sandbox"
    )

    options.add_argument(
        "--disable-dev-shm-usage"
    )

    options.add_argument(
        "--lang=en-US"
    )

    options.add_experimental_option(
        "excludeSwitches",
        ["enable-automation"]
    )

    options.add_experimental_option(
        "useAutomationExtension",
        False
    )

    driver = webdriver.Chrome(
        options=options
    )

    # Set page load timeout
    driver.set_page_load_timeout(60)

    return driver


# ============================================================
# SCRAPE ONE PAGE
# ============================================================

def scrape_area(driver, area, slug):

    url = f"{BASE_URL}/{slug}"

    print()
    print("=" * 75)
    print(f"AREA : {area}")
    print(f"URL  : {url}")
    print("=" * 75)

    try:

        # ----------------------------------------------------
        # Open page
        # ----------------------------------------------------

        driver.get(url)

        print("Page opened.")

        # ----------------------------------------------------
        # Wait for page body
        # ----------------------------------------------------

        WebDriverWait(
            driver,
            WAIT_TIME
        ).until(
            EC.presence_of_element_located(
                (By.TAG_NAME, "body")
            )
        )

        # ----------------------------------------------------
        # Wait for tables
        # ----------------------------------------------------

        try:

            WebDriverWait(
                driver,
                WAIT_TIME
            ).until(
                lambda d: len(
                    d.find_elements(
                        By.TAG_NAME,
                        "table"
                    )
                ) >= 2
            )

        except Exception:

            print(
                "WARNING: Two tables did not appear "
                "within the wait time."
            )

        # ----------------------------------------------------
        # Small additional wait
        # ----------------------------------------------------

        time.sleep(2)

        # ----------------------------------------------------
        # Get rendered HTML
        # ----------------------------------------------------

        html = driver.page_source

        # ----------------------------------------------------
        # Save HTML for debugging
        # ----------------------------------------------------

        debug_dir = "IGL_Debug_HTML"

        os.makedirs(
            debug_dir,
            exist_ok=True
        )

        debug_file = os.path.join(
            debug_dir,
            f"{slug}.html"
        )

        with open(
            debug_file,
            "w",
            encoding="utf-8"
        ) as f:

            f.write(html)

        # ----------------------------------------------------
        # Parse HTML
        # ----------------------------------------------------

        soup = BeautifulSoup(
            html,
            "html.parser"
        )

        tables = soup.find_all("table")

        print(
            f"Tables detected: {len(tables)}"
        )

        # ----------------------------------------------------
        # Check tables
        # ----------------------------------------------------

        if len(tables) == 0:

            print(
                "ERROR: No tables found."
            )

            return (
                pd.DataFrame(),
                "No tables found"
            )

        # ----------------------------------------------------
        # Extract tables
        # ----------------------------------------------------

        page_data = []

        # We only need the first two tables
        # as requested.

        for table_number, table in enumerate(
            tables[:2],
            start=1
        ):

            rows = []

            # ------------------------------------------------
            # Header
            # ------------------------------------------------

            header_row = table.find("tr")

            if not header_row:

                continue

            headers = [
                clean_text(
                    cell.get_text(
                        " ",
                        strip=True
                    )
                )
                for cell in header_row.find_all(
                    ["th", "td"]
                )
            ]

            print(
                f"Table {table_number} headers: "
                f"{headers}"
            )

            # ------------------------------------------------
            # Data rows
            # ------------------------------------------------

            for tr in table.find_all("tr")[1:]:

                cells = tr.find_all(
                    ["td", "th"]
                )

                values = [
                    clean_text(
                        cell.get_text(
                            " ",
                            strip=True
                        )
                    )
                    for cell in cells
                ]

                # Ignore empty rows
                if not values:
                    continue

                if all(
                    value == ""
                    for value in values
                ):
                    continue

                rows.append(values)

            if not rows:

                print(
                    f"Table {table_number}: "
                    f"No data rows"
                )

                continue

            # ------------------------------------------------
            # Make dataframe
            # ------------------------------------------------

            # Handle mismatched column counts
            fixed_rows = []

            for row in rows:

                if len(row) < len(headers):

                    row = row + [
                        ""
                    ] * (
                        len(headers) - len(row)
                    )

                elif len(row) > len(headers):

                    row = row[
                        :len(headers)
                    ]

                fixed_rows.append(row)

            df = pd.DataFrame(
                fixed_rows,
                columns=headers
            )

            # ------------------------------------------------
            # Identify table type
            # ------------------------------------------------

            if (
                "Phone Number"
                in headers
            ):

                table_type = "CNG Station"

            elif (
                "OMC"
                in headers
            ):

                table_type = "OMC Outlet"

            else:

                table_type = (
                    f"Table {table_number}"
                )

            # ------------------------------------------------
            # Add metadata
            # ------------------------------------------------

            df.insert(
                0,
                "Area",
                area
            )

            df.insert(
                1,
                "Table_Type",
                table_type
            )

            df["Source_URL"] = url

            page_data.append(df)

            print(
                f"Table {table_number}: "
                f"{len(df)} records"
            )

        # ----------------------------------------------------
        # Combine page tables
        # ----------------------------------------------------

        if not page_data:

            return (
                pd.DataFrame(),
                "Tables found but no records"
            )

        result = pd.concat(
            page_data,
            ignore_index=True,
            sort=False
        )

        print(
            f"TOTAL RECORDS FOR {area}: "
            f"{len(result)}"
        )

        return (
            result,
            "Success"
        )

    except Exception as e:

        print(
            f"ERROR scraping {area}: {e}"
        )

        # Save screenshot
        try:

            screenshot_dir = (
                "IGL_Debug_Screenshots"
            )

            os.makedirs(
                screenshot_dir,
                exist_ok=True
            )

            screenshot_file = os.path.join(
                screenshot_dir,
                f"{slug}.png"
            )

            driver.save_screenshot(
                screenshot_file
            )

            print(
                f"Screenshot saved: "
                f"{screenshot_file}"
            )

        except Exception:
            pass

        return (
            pd.DataFrame(),
            str(e)
        )


# ============================================================
# MAIN
# ============================================================

def main():

    start_time = time.time()

    print()
    print("#" * 75)
    print("IGL CNG STATION SCRAPER")
    print("#" * 75)

    print()
    print(
        f"Total areas: {len(AREAS)}"
    )

    # --------------------------------------------------------
    # Create Chrome
    # --------------------------------------------------------

    print()
    print(
        "Starting Chrome..."
    )

    driver = create_driver()

    all_data = []

    scrape_log = []

    try:

        # ====================================================
        # PROCESS EACH AREA
        # ====================================================

        for index, (
            area,
            slug
        ) in enumerate(
            AREAS,
            start=1
        ):

            print()
            print(
                f"[{index}/{len(AREAS)}] "
                f"Processing {area}"
            )

            df, status = scrape_area(
                driver,
                area,
                slug
            )

            if not df.empty:

                all_data.append(df)

                scrape_log.append({
                    "Area": area,
                    "URL": (
                        f"{BASE_URL}/{slug}"
                    ),
                    "Status": "Success",
                    "Records": len(df),
                    "Error": ""
                })

            else:

                scrape_log.append({
                    "Area": area,
                    "URL": (
                        f"{BASE_URL}/{slug}"
                    ),
                    "Status": "Failed / Empty",
                    "Records": 0,
                    "Error": status
                })

            time.sleep(
                PAGE_DELAY
            )

    finally:

        print()
        print(
            "Closing Chrome..."
        )

        driver.quit()

    # ========================================================
    # CHECK RESULT
    # ========================================================

    if not all_data:

        print()
        print("#" * 75)
        print(
            "NO DATA WAS SCRAPED"
        )
        print("#" * 75)

        print()
        print(
            "Check the files inside:"
        )

        print(
            "IGL_Debug_HTML"
        )

        print(
            "and"
        )

        print(
            "IGL_Debug_Screenshots"
        )

        return

    # ========================================================
    # COMBINE DATA
    # ========================================================

    print()
    print(
        "=" * 75
    )

    print(
        "Combining all area data..."
    )

    final_df = pd.concat(
        all_data,
        ignore_index=True,
        sort=False
    )

    # ========================================================
    # REMOVE EMPTY STATION NAMES
    # ========================================================

    if "Station Name" in final_df.columns:

        final_df = final_df[
            final_df[
                "Station Name"
            ].astype(str).str.strip() != ""
        ]

    # ========================================================
    # REMOVE DUPLICATES
    # ========================================================

    duplicate_columns = [
        col
        for col in [
            "Area",
            "Table_Type",
            "Station Name",
            "Address"
        ]
        if col in final_df.columns
    ]

    before_duplicates = len(
        final_df
    )

    final_df = final_df.drop_duplicates(
        subset=duplicate_columns
    )

    duplicates_removed = (
        before_duplicates
        - len(final_df)
    )

    # ========================================================
    # COLUMN ORDER
    # ========================================================

    preferred_columns = [
        "Area",
        "Table_Type",
        "S.No",
        "Station Name",
        "Address",
        "Phone Number",
        "OMC",
        "Source_URL"
    ]

    existing_columns = [
        col
        for col in preferred_columns
        if col in final_df.columns
    ]

    remaining_columns = [
        col
        for col in final_df.columns
        if col not in existing_columns
    ]

    final_df = final_df[
        existing_columns
        + remaining_columns
    ]

    # ========================================================
    # LOG
    # ========================================================

    log_df = pd.DataFrame(
        scrape_log
    )

    # ========================================================
    # SUMMARY
    # ========================================================

    successful_areas = len(
        log_df[
            log_df["Status"]
            == "Success"
        ]
    )

    failed_areas = len(
        log_df[
            log_df["Status"]
            != "Success"
        ]
    )

    summary_df = pd.DataFrame({
        "Metric": [
            "Source Website",
            "Total Areas",
            "Successful Areas",
            "Failed / Empty Areas",
            "Total Records",
            "Duplicates Removed"
        ],

        "Value": [
            BASE_URL,
            len(AREAS),
            successful_areas,
            failed_areas,
            len(final_df),
            duplicates_removed
        ]
    })

    # ========================================================
    # SAVE EXCEL
    # ========================================================

    print()
    print(
        "=" * 75
    )

    print(
        "Saving Excel..."
    )

    with pd.ExcelWriter(
        OUTPUT_FILE,
        engine="openpyxl"
    ) as writer:

        final_df.to_excel(
            writer,
            sheet_name="All_CNG_Stations",
            index=False
        )

        log_df.to_excel(
            writer,
            sheet_name="Scrape_Log",
            index=False
        )

        summary_df.to_excel(
            writer,
            sheet_name="Summary",
            index=False
        )

    # ========================================================
    # FORMAT EXCEL
    # ========================================================

    try:

        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        wb = load_workbook(
            OUTPUT_FILE
        )

        for ws in wb.worksheets:

            # Freeze header
            ws.freeze_panes = "A2"

            # Filter
            ws.auto_filter.ref = (
                ws.dimensions
            )

            # Column width
            for column_cells in ws.columns:

                max_length = 0

                column_letter = (
                    get_column_letter(
                        column_cells[0].column
                    )
                )

                for cell in column_cells:

                    try:

                        length = len(
                            str(cell.value)
                        )

                        max_length = max(
                            max_length,
                            length
                        )

                    except:
                        pass

                ws.column_dimensions[
                    column_letter
                ].width = min(
                    max(
                        max_length + 2,
                        12
                    ),
                    60
                )

        wb.save(
            OUTPUT_FILE
        )

    except Exception as e:

        print(
            f"Excel formatting warning: {e}"
        )

    # ========================================================
    # FINAL SUMMARY
    # ========================================================

    elapsed = (
        time.time()
        - start_time
    )

    print()
    print("#" * 75)
    print("SCRAPING COMPLETED")
    print("#" * 75)

    print(
        f"Total areas       : "
        f"{len(AREAS)}"
    )

    print(
        f"Successful areas  : "
        f"{successful_areas}"
    )

    print(
        f"Failed / Empty    : "
        f"{failed_areas}"
    )

    print(
        f"Total records     : "
        f"{len(final_df)}"
    )

    print(
        f"Duplicates removed: "
        f"{duplicates_removed}"
    )

    print(
        f"Excel file        : "
        f"{os.path.abspath(OUTPUT_FILE)}"
    )

    print(
        f"Time taken        : "
        f"{elapsed:.2f} seconds"
    )

    # ========================================================
    # FAILED AREAS
    # ========================================================

    if failed_areas > 0:

        print()
        print(
            "=" * 75
        )

        print(
            "FAILED AREAS"
        )

        print(
            "=" * 75
        )

        failed_df = log_df[
            log_df["Status"]
            != "Success"
        ]

        for _, row in failed_df.iterrows():

            print(
                f"{row['Area']} -> "
                f"{row['Error']}"
            )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    main()